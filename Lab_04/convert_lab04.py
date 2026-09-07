"""Convert 23MID0021_Lab04_Report.md to .docx and .pdf."""

import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import win32com.client

BASE = r'c:\Users\I768951\OneDrive - SAP SE\SAP\Exam'
MD_PATH   = os.path.join(BASE, '23MID0021_Lab04_Report.md')
DOCX_PATH = os.path.join(BASE, '23MID0021_Lab04_Report.docx')
PDF_PATH  = os.path.join(BASE, '23MID0021_Lab04_Report.pdf')


# ── inline markdown → paragraph runs ─────────────────────────────────────────

def add_inline(para, text):
    """Render inline **bold**, *italic*, `code`, $math$ into a paragraph."""
    token = re.split(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\$[^$\n]+?\$)', text)
    for t in token:
        if not t:
            continue
        if t.startswith('**') and t.endswith('**'):
            r = para.add_run(t[2:-2]); r.bold = True
        elif t.startswith('*') and t.endswith('*'):
            r = para.add_run(t[1:-1]); r.italic = True
        elif t.startswith('`') and t.endswith('`'):
            r = para.add_run(t[1:-1])
            r.font.name = 'Courier New'; r.font.size = Pt(9)
        elif t.startswith('$') and t.endswith('$'):
            r = para.add_run(t.strip('$'))
            r.italic = True; r.font.name = 'Cambria Math'
        else:
            para.add_run(t)


def strip_inline(text):
    """Plain text without markdown markers."""
    text = re.sub(r'\*\*([^*]+?)\*\*', r'\1', text)
    text = re.sub(r'\*([^*]+?)\*',     r'\1', text)
    text = re.sub(r'`([^`]+?)`',       r'\1', text)
    text = re.sub(r'\$([^$\n]+?)\$',   r'\1', text)
    return text


# ── table helper ─────────────────────────────────────────────────────────────

def shade_cell(cell, fill='D9E1F2'):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  fill)
    tcPr.append(shd)


def add_table(doc, table_lines):
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        # skip separator rows (e.g. | :--- | :---: |)
        if all(re.fullmatch(r'[-: ]+', c) for c in cells if c):
            continue
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    tbl   = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = 'Table Grid'
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = tbl.rows[ri].cells[ci]
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            val  = row[ci] if ci < len(row) else ''
            if ri == 0:
                r = para.add_run(strip_inline(val))
                r.bold = True
                shade_cell(cell)
            else:
                add_inline(para, val)
            # compact row height
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after  = Pt(2)


# ── code block helper ─────────────────────────────────────────────────────────

def add_code_block(doc, code_text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after  = Pt(4)
    pPr  = para._p.get_or_add_pPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'F2F2F2')
    pPr.append(shd)
    run = para.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)


# ── main conversion ───────────────────────────────────────────────────────────

def md_to_docx(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc  = Document()
    sect = doc.sections[0]
    sect.top_margin    = Cm(2.5)
    sect.bottom_margin = Cm(2.5)
    sect.left_margin   = Cm(3.0)
    sect.right_margin  = Cm(2.5)

    # Default body font
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ── Headings ──────────────────────────────────────────────────────────
        if re.match(r'^#{1,4} ', line):
            level = len(re.match(r'^(#+)', line).group(1))
            text  = line.lstrip('#').strip()
            h = doc.add_heading(text, level=min(level, 4))
            if level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue

        # ── Horizontal rule ───────────────────────────────────────────────────
        if stripped == '---':
            p   = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bot  = OxmlElement('w:bottom')
            bot.set(qn('w:val'),   'single')
            bot.set(qn('w:sz'),    '6')
            bot.set(qn('w:space'), '1')
            bot.set(qn('w:color'), 'AAAAAA')
            pBdr.append(bot)
            pPr.append(pBdr)
            i += 1; continue

        # ── Fenced code block ─────────────────────────────────────────────────
        if stripped.startswith('```'):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            add_code_block(doc, '\n'.join(code))
            i += 1; continue

        # ── Display math ($$…$$) ──────────────────────────────────────────────
        if stripped.startswith('$$'):
            math_parts = [stripped[2:]]
            if not stripped.endswith('$$') or stripped == '$$':
                i += 1
                while i < len(lines) and not lines[i].strip().endswith('$$'):
                    math_parts.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    math_parts.append(lines[i].strip().rstrip('$'))
            math_text = ' '.join(p for p in math_parts if p).strip('$ ')
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(math_text)
            r.italic = True; r.font.name = 'Cambria Math'; r.font.size = Pt(11)
            i += 1; continue

        # ── Table ─────────────────────────────────────────────────────────────
        if stripped.startswith('|'):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl_lines.append(lines[i])
                i += 1
            add_table(doc, tbl_lines)
            continue

        # ── Bullet list ───────────────────────────────────────────────────────
        if re.match(r'^[-*] ', line):
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, line[2:].strip())
            i += 1; continue

        # ── Numbered list ─────────────────────────────────────────────────────
        if re.match(r'^\d+\. ', line):
            p = doc.add_paragraph(style='List Number')
            add_inline(p, re.sub(r'^\d+\. ', '', line))
            i += 1; continue

        # ── Figure image line (skip image, keep alt-text as italic caption) ──
        if stripped.startswith('!['):
            m = re.match(r'!\[([^\]]*)\]', stripped)
            if m:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(f'[Figure: {m.group(1)}]')
                r.italic = True
                r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            i += 1; continue

        # ── Italic-only line (figure captions: *text*) ────────────────────────
        if stripped.startswith('*') and stripped.endswith('*') \
                and not stripped.startswith('**') and len(stripped) > 2:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(stripped[1:-1])
            r.italic = True; r.font.size = Pt(10)
            p.paragraph_format.space_before = Pt(0)
            i += 1; continue

        # ── Empty line ────────────────────────────────────────────────────────
        if stripped == '':
            i += 1; continue

        # ── Normal paragraph ──────────────────────────────────────────────────
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f'Word saved: {docx_path}')


def docx_to_pdf(docx_path, pdf_path):
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(docx_path)
        doc.SaveAs(pdf_path, FileFormat=17)   # 17 = wdFormatPDF
        doc.Close()
        print(f'PDF  saved: {pdf_path}')
    finally:
        word.Quit()


if __name__ == '__main__':
    md_to_docx(MD_PATH, DOCX_PATH)
    docx_to_pdf(DOCX_PATH, PDF_PATH)
    print('Done.')
