"""Build 23MID0021_Lab05_Report.docx from markdown and export PDF.
Follows the same structure as convert_lab04.py but embeds figures directly.
"""
import os, re
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import win32com.client

BASE      = r'c:\Users\I768951\OneDrive - SAP SE\SAP\Exam'
MD_PATH   = os.path.join(BASE, '23MID0021_Lab05_Report.md')
DOCX_PATH = os.path.join(BASE, '23MID0021_Lab05_Report.docx')
PDF_PATH  = os.path.join(BASE, '23MID0021_Lab05_Report.pdf')
FIG_DIR   = os.path.join(BASE, 'lab05_figs')

# Map markdown image ref basename → actual figure file
FIG_MAP = {
    'fig1_class_dist.png':          os.path.join(FIG_DIR, 'fig1_class_dist.png'),
    'fig2_tweet_length.png':        os.path.join(FIG_DIR, 'fig2_tweet_length.png'),
    'fig3_top_terms.png':           os.path.join(FIG_DIR, 'fig3_top_terms.png'),
    'fig4_airline_dist.png':        os.path.join(FIG_DIR, 'fig4_airline_dist.png'),
    'fig5_cv_comparison.png':       os.path.join(FIG_DIR, 'fig5_cv_comparison.png'),
    'fig6_confusion_matrix.png':    os.path.join(FIG_DIR, 'fig6_confusion_matrix.png'),
    'fig7_per_class_f1.png':        os.path.join(FIG_DIR, 'fig7_per_class_f1.png'),
    'fig8_vader_vs_best.png':       os.path.join(FIG_DIR, 'fig8_vader_vs_best.png'),
    'fig9_entity_analysis.png':     os.path.join(FIG_DIR, 'fig9_entity_analysis.png'),
    'fig10_error_analysis.png':     os.path.join(FIG_DIR, 'fig10_error_analysis.png'),
    'fig11_discriminative_terms.png': os.path.join(FIG_DIR, 'fig11_discriminative_terms.png'),
    'fig12_learning_curve.png':     os.path.join(FIG_DIR, 'fig12_learning_curve.png'),
    'fig13_pr_curves.png':          os.path.join(FIG_DIR, 'fig13_pr_curves.png'),
    'fig14_feature_heatmap.png':    os.path.join(FIG_DIR, 'fig14_feature_heatmap.png'),
    'fig15_confidence_dist.png':    os.path.join(FIG_DIR, 'fig15_confidence_dist.png'),
    'fig16_airline_grouped_bar.png': os.path.join(FIG_DIR, 'fig16_airline_grouped_bar.png'),
    'fig17_cv_detailed.png':        os.path.join(FIG_DIR, 'fig17_cv_detailed.png'),
}

# Figures that need wider width (6.2"); rest use 5.5"
WIDE_FIGS = {
    'fig6_confusion_matrix.png', 'fig8_vader_vs_best.png',
    'fig9_entity_analysis.png',  'fig10_error_analysis.png',
    'fig11_discriminative_terms.png', 'fig14_feature_heatmap.png',
    'fig15_confidence_dist.png', 'fig16_airline_grouped_bar.png',
    'fig17_cv_detailed.png', 'fig7_per_class_f1.png',
    'fig13_pr_curves.png',
}


# ── inline markdown → paragraph runs ─────────────────────────────────────────

def add_inline(para, text):
    token = re.split(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`|\$[^$\n]+?\$)', text)
    for t in token:
        if not t:
            continue
        if t.startswith('**') and t.endswith('**'):
            r = para.add_run(t[2:-2]); r.bold = True
        elif t.startswith('*') and t.endswith('*'):
            r = para.add_run(t[1:-1]); r.italic = True
        elif t.startswith('`') and t.endswith('`'):
            r = para.add_run(t[1:-1])
            r.font.name = 'Courier New'; r.font.size = Pt(9)
        elif t.startswith('$') and t.endswith('$'):
            r = para.add_run(t.strip('$'))
            r.italic = True; r.font.name = 'Cambria Math'
        else:
            para.add_run(t)


def strip_inline(text):
    text = re.sub(r'\*\*([^*]+?)\*\*', r'\1', text)
    text = re.sub(r'\*([^*]+?)\*',     r'\1', text)
    text = re.sub(r'`([^`]+?)`',       r'\1', text)
    text = re.sub(r'\$([^$\n]+?)\$',   r'\1', text)
    return text


# ── helpers ───────────────────────────────────────────────────────────────────

def shade_cell(cell, fill='D9E1F2'):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  fill)
    tcPr.append(shd)


def add_table(doc, table_lines):
    rows = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r'[-: ]+', c) for c in cells if c):
            continue
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    tbl   = doc.add_table(rows=len(rows), cols=ncols)
    tbl.style = 'Table Grid'
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = tbl.rows[ri].cells[ci]
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            val  = row[ci] if ci < len(row) else ''
            if ri == 0:
                r = para.add_run(strip_inline(val))
                r.bold = True
                shade_cell(cell)
            else:
                add_inline(para, val)
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after  = Pt(2)


def add_code_block(doc, code_text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after  = Pt(4)
    para.paragraph_format.left_indent  = Cm(0.3)
    pPr  = para._p.get_or_add_pPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  'F2F2F2')
    pPr.append(shd)
    run = para.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)


def add_figure(doc, filename):
    fig_path = FIG_MAP.get(filename)
    if not fig_path or not os.path.exists(fig_path):
        print(f'  MISSING: {filename}')
        return
    width = Inches(6.2) if filename in WIDE_FIGS else Inches(5.5)
    para  = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after  = Pt(2)
    run = para.add_run()
    run.add_picture(fig_path, width=width)
    print(f'  Embedded: {filename}')


def add_header_footer(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    hdr = section.header
    hdr.is_linked_to_previous = False
    hp  = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run('MDI3003 - Advanced Predictive Analytics  |  Laboratory Report')
    r.font.name = 'Calibri'; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    ftr = section.footer
    ftr.is_linked_to_previous = False
    fp  = ftr.paragraphs[0] if ftr.paragraphs else ftr.add_paragraph()
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('Geetha Priya S  |  23MID0021  |  Lab 05  |  Page ')
    r.font.name = 'Calibri'; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    # Page number field -- each w:fldChar must be inside a w:r
    for fld_type in ('begin', None, 'end'):
        rr = OxmlElement('w:r')
        if fld_type is None:
            instr = OxmlElement('w:instrText')
            instr.text = ' PAGE '
            rr.append(instr)
        else:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), fld_type)
            rr.append(fc)
        fp._p.append(rr)


# ── main conversion ───────────────────────────────────────────────────────────

def md_to_docx(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc  = Document()
    sect = doc.sections[0]
    sect.top_margin    = Cm(2.5)
    sect.bottom_margin = Cm(2.5)
    sect.left_margin   = Cm(3.0)
    sect.right_margin  = Cm(2.5)

    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    add_header_footer(doc)

    i = 0
    while i < len(lines):
        line     = lines[i]
        stripped = line.strip()

        # ── Headings ──────────────────────────────────────────────────────────
        if re.match(r'^#{1,4} ', line):
            level = len(re.match(r'^(#+)', line).group(1))
            text  = line.lstrip('#').strip()
            h = doc.add_heading(strip_inline(text), level=min(level, 4))
            if level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue

        # ── Horizontal rule ───────────────────────────────────────────────────
        if stripped == '---':
            p    = doc.add_paragraph()
            pPr  = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bot  = OxmlElement('w:bottom')
            bot.set(qn('w:val'),   'single')
            bot.set(qn('w:sz'),    '6')
            bot.set(qn('w:space'), '1')
            bot.set(qn('w:color'), 'AAAAAA')
            pBdr.append(bot)
            pPr.append(pBdr)
            i += 1; continue

        # ── Fenced code block ─────────────────────────────────────────────────
        if stripped.startswith('```'):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            add_code_block(doc, '\n'.join(code))
            i += 1; continue

        # ── Display math ($$...$$) ────────────────────────────────────────────
        if stripped.startswith('$$'):
            math_parts = [stripped[2:]]
            if not stripped.endswith('$$') or stripped == '$$':
                i += 1
                while i < len(lines) and not lines[i].strip().endswith('$$'):
                    math_parts.append(lines[i].strip())
                    i += 1
                if i < len(lines):
                    math_parts.append(lines[i].strip().rstrip('$'))
            math_text = ' '.join(p for p in math_parts if p).strip('$ ')
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(math_text)
            r.italic = True; r.font.name = 'Cambria Math'; r.font.size = Pt(11)
            i += 1; continue

        # ── Table ─────────────────────────────────────────────────────────────
        if stripped.startswith('|'):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                tbl_lines.append(lines[i])
                i += 1
            add_table(doc, tbl_lines)
            continue

        # ── Figure image (embed directly) ─────────────────────────────────────
        if stripped.startswith('!['):
            m = re.match(r'!\[[^\]]*\]\(figures/([^)]+)\)', stripped)
            if m:
                add_figure(doc, m.group(1))
            i += 1; continue

        # ── Italic-only line (figure captions: *text*) ────────────────────────
        if stripped.startswith('*') and stripped.endswith('*') \
                and not stripped.startswith('**') and len(stripped) > 2:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(stripped[1:-1])
            r.italic = True; r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(8)
            i += 1; continue

        # ── Bullet list ───────────────────────────────────────────────────────
        if re.match(r'^[-*] ', line):
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, line[2:].strip())
            i += 1; continue

        # ── Numbered list ─────────────────────────────────────────────────────
        if re.match(r'^\d+\. ', line):
            p = doc.add_paragraph(style='List Number')
            add_inline(p, re.sub(r'^\d+\. ', '', line))
            i += 1; continue

        # ── Empty line ────────────────────────────────────────────────────────
        if stripped == '':
            i += 1; continue

        # ── Normal paragraph ──────────────────────────────────────────────────
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f'\nWord saved: {docx_path}')


def docx_to_pdf(docx_path, pdf_path):
    print('Exporting PDF via Word COM...')
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    import shutil
    tmp = r'c:\Temp\Lab05_export.docx'
    os.makedirs(r'c:\Temp', exist_ok=True)
    shutil.copy2(docx_path, tmp)
    try:
        doc = word.Documents.Open(tmp)
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
        doc.Close()
        print(f'PDF  saved: {pdf_path}')
    finally:
        word.Quit()


if __name__ == '__main__':
    print('=== Lab 05 Report Builder ===\n')
    print('Checking figures...')
    ok = True
    for fname, fpath in FIG_MAP.items():
        status = 'OK' if os.path.exists(fpath) else 'MISSING'
        if status == 'MISSING': ok = False
        print(f'  [{status}] {fname}')
    if not ok:
        print('\nWARNING: some figures are missing.')
    print()
    md_to_docx(MD_PATH, DOCX_PATH)
    print()
    docx_to_pdf(DOCX_PATH, PDF_PATH)
    print('\nDone.')
